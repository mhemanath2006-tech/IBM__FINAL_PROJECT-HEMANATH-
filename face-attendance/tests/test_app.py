"""End-to-end tests with a FAKE face engine.

The fake engine reads identity / head-yaw / face-count out of the first pixel of a
PNG, so these tests prove the app logic (roles, sessions, windows, challenge flow,
reports) but NOT real face recognition accuracy. Real faces must be tested on the
deployed site.
"""
import base64
import io
import os
import sys
import tempfile
import unittest

import cv2
import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
_tmp = tempfile.mkdtemp()
os.environ["SQLITE_PATH"] = os.path.join(_tmp, "t.db")
os.environ["ADMIN_PASSWORD"] = "adminpass"
os.environ["SECRET_KEY"] = "test"
os.environ.pop("DATABASE_URL", None)

import face_engine as fe  # noqa: E402
from app import create_app  # noqa: E402
from db import _connect  # noqa: E402


def png(identity, yaw=128, faces=1):
    img = np.zeros((8, 8, 3), np.uint8)
    img[:, :] = (identity, yaw, faces)  # B, G, R
    ok, buf = cv2.imencode(".png", img)
    return "data:image/png;base64," + base64.b64encode(buf.tobytes()).decode()


class FakeEngine:
    def analyze(self, img):
        b, gch, r = (int(x) for x in img[0, 0])
        if r == 0:
            return []
        emb = np.random.RandomState(b).randn(512).astype(np.float32)
        emb += np.random.RandomState(gch + 1000).randn(512).astype(np.float32) * 0.02  # tiny pose noise
        emb /= np.linalg.norm(emb)
        kps = np.array([[100, 100], [160, 100], [130 + (gch - 128) * 0.5, 130], [110, 160], [150, 160]], np.float32)
        return [{"embedding": emb, "kps": kps, "score": 0.99, "bbox": np.array([50, 50, 250, 250], np.float32)}
                for _ in range(r)]


LEFT, RIGHT, STRAIGHT = 152, 104, 128  # yaw codes -> +0.2 / -0.2 / 0.0


class AppTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fe.set_engine(FakeEngine())
        cls.app = create_app(testing=True)

    def client(self, user=None, pw=None):
        c = self.app.test_client()
        c.get("/login")
        if user:
            r = c.post("/login", data={"username": user, "password": pw, "csrf": self.csrf(c)})
            self.assertEqual(r.status_code, 302, "login failed for " + user)
        return c

    @staticmethod
    def csrf(c):
        with c.session_transaction() as s:
            s.setdefault("csrf", "tok")
            return s["csrf"]

    def post(self, c, url, **data):
        data["csrf"] = self.csrf(c)
        return c.post(url, data=data)

    def jpost(self, c, url, body):
        return c.post(url, json=body, headers={"X-CSRF-Token": self.csrf(c)})

    def make_user(self, admin, role, name, roll, section, username=None, pw="pass1234"):
        self.post(admin, "/admin/users", role=role, name=name, roll_no=roll, section=section,
                  username=username or roll, password=pw)

    def uid(self, username):
        return _connect().one("SELECT id FROM users WHERE username = ?", (username,))["id"]

    def enroll(self, admin, username, identity):
        frames = [png(identity, y) for y in (STRAIGHT, LEFT, RIGHT, STRAIGHT, STRAIGHT)]
        return self.jpost(admin, f"/api/admin/enroll/{self.uid(username)}", {"frames": frames, "consent": True})

    def start_session(self, teacher, section="CSE-A", open_m=30, late_m=10):
        r = self.post(teacher, "/teacher/sessions", section=section, open_minutes=open_m, late_minutes=late_m)
        self.assertEqual(r.status_code, 302)
        return int(r.headers["Location"].rstrip("/").split("/")[-1])

    def scan(self, student, identity, neutral_yaw=STRAIGHT, turn_yaws=(STRAIGHT, 140, LEFT, LEFT, 150)):
        ch = self.jpost(student, "/api/scan/challenge", {}).get_json()
        if not ch["ok"]:
            return ch
        want = ch["direction"]
        turns = [png(identity, y if want == "left" else 256 - y) for y in turn_yaws]
        return self.jpost(student, "/api/scan/submit",
                          {"nonce": ch["nonce"], "neutral": png(identity, neutral_yaw), "turns": turns}).get_json()

    # ---------------------------------------------------------------- tests
    def test_full_flow(self):
        admin = self.client("admin", "adminpass")
        self.make_user(admin, "teacher", "Mrs Rao", "", "", username="rao")
        for i, (name, roll) in enumerate([("Asha", "A1"), ("Vik", "A2"), ("Meera", "A3")], start=1):
            self.make_user(admin, "student", name, roll, "FLOW")
            self.assertTrue(self.enroll(admin, roll, 20 + i).get_json()["ok"])
        teacher = self.client("rao", "pass1234")
        sid = self.start_session(teacher, "FLOW")

        asha = self.client("A1", "pass1234")
        # Correct person, real head turn -> present
        r = self.scan(asha, 21)
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["status"], "present")
        # Already marked -> no second scan
        self.assertFalse(self.scan(asha, 21)["ok"])

        vik = self.client("A2", "pass1234")
        # Someone else's face on Vik's account is rejected
        r = self.scan(vik, 99)
        self.assertFalse(r["ok"])
        self.assertIn("did not match", r["error"])
        # Photo (no head movement) is rejected by liveness
        r = self.scan(vik, 22, turn_yaws=(STRAIGHT,) * 5)
        self.assertFalse(r["ok"])
        self.assertIn("Liveness", r["error"])
        # Turning the wrong way is rejected
        ch = self.jpost(vik, "/api/scan/challenge", {}).get_json()
        wrong = "right" if ch["direction"] == "left" else "left"
        yaw = RIGHT if ch["direction"] == "left" else LEFT
        r = self.jpost(vik, "/api/scan/submit", {"nonce": ch["nonce"], "neutral": png(22),
                                                 "turns": [png(22, yaw)] * 5}).get_json()
        self.assertFalse(r["ok"], wrong)
        # Used challenge cannot be replayed
        r = self.jpost(vik, "/api/scan/submit", {"nonce": ch["nonce"], "neutral": png(22),
                                                 "turns": [png(22, yaw)] * 5}).get_json()
        self.assertFalse(r["ok"])
        self.assertIn("expired", r["error"])
        # Two people in frame rejected
        ch = self.jpost(vik, "/api/scan/challenge", {}).get_json()
        r = self.jpost(vik, "/api/scan/submit", {"nonce": ch["nonce"], "neutral": png(22, faces=2),
                                                 "turns": [png(22)] * 5}).get_json()
        self.assertFalse(r["ok"])
        # Vik succeeds after a proper scan
        self.assertTrue(self.scan(vik, 22)["ok"])

        # Late: move the late boundary into the past, Meera scans -> late
        _connect().run("UPDATE class_sessions SET late_after = '2000-01-01 00:00:00' WHERE id = ?", (sid,))
        meera = self.client("A3", "pass1234")
        self.assertEqual(self.scan(meera, 23)["status"], "late")

        page = teacher.get(f"/teacher/sessions/{sid}").get_data(as_text=True)
        self.assertIn("Asha", page)

        # Teacher manual override, then close and verify scanning stops
        self.post(teacher, f"/teacher/sessions/{sid}/mark", student_id=self.uid("A3"), status="absent")
        row = _connect().one("SELECT status, method FROM attendance WHERE student_id = ? AND session_id = ?",
                             (self.uid("A3"), sid))
        self.assertEqual((row["status"], row["method"]), ("absent", "manual"))
        self.post(teacher, f"/teacher/sessions/{sid}/close")
        self.assertFalse(self.jpost(asha, "/api/scan/challenge", {}).get_json()["ok"])

        # Reports + Excel
        rep = teacher.get("/teacher/reports?section=FLOW").get_data(as_text=True)
        self.assertIn("Asha", rep)
        x = teacher.get("/teacher/export.xlsx?section=FLOW")
        self.assertEqual(x.status_code, 200)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(x.data))
        rows = {r[1]: r for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["Asha"][4], 1)       # present
        self.assertEqual(rows["Vik"][4], 1)
        self.assertEqual(rows["Meera"][7], 0.0)    # marked absent manually -> 0%

        # Student sees own dashboard
        self.assertIn("Present", asha.get("/student").get_data(as_text=True))

    def test_roles_and_csrf(self):
        admin = self.client("admin", "adminpass")
        self.make_user(admin, "student", "Role Kid", "R1", "ROLES")
        s = self.client("R1", "pass1234")
        self.assertEqual(s.get("/admin").status_code, 403)
        self.assertEqual(s.get("/teacher").status_code, 403)
        self.assertEqual(self.client().get("/student").status_code, 302)  # anonymous -> login
        self.assertEqual(s.post("/logout", data={}).status_code, 400)       # missing CSRF
        r = self.client().post("/api/scan/challenge", json={})
        self.assertEqual(r.status_code, 400)

    def test_enroll_rejects_mixed_people_and_needs_consent(self):
        admin = self.client("admin", "adminpass")
        self.make_user(admin, "student", "Mix Kid", "M1", "MIX")
        uid = self.uid("M1")
        frames = [png(40), png(41), png(42), png(43)]
        r = self.jpost(admin, f"/api/admin/enroll/{uid}", {"frames": frames, "consent": True})
        self.assertFalse(r.get_json()["ok"])
        r = self.jpost(admin, f"/api/admin/enroll/{uid}", {"frames": [png(40)] * 4, "consent": False})
        self.assertEqual(r.status_code, 400)
        r = self.jpost(admin, f"/api/admin/enroll/{uid}", {"frames": [png(40, faces=0)] * 4, "consent": True})
        self.assertFalse(r.get_json()["ok"])

    def test_no_session_and_unregistered_face(self):
        admin = self.client("admin", "adminpass")
        self.make_user(admin, "student", "Nofa", "N1", "NONE")
        s = self.client("N1", "pass1234")
        self.assertIn("not registered", self.jpost(s, "/api/scan/challenge", {}).get_json()["error"])
        self.enroll(admin, "N1", 60)
        self.assertIn("not open", self.jpost(s, "/api/scan/challenge", {}).get_json()["error"])

    def test_every_page_renders(self):
        admin = self.client("admin", "adminpass")
        self.make_user(admin, "student", "Page Kid", "P1", "PAGES")
        self.make_user(admin, "teacher", "Page Teacher", "", "", username="pt")
        uid = self.uid("P1")
        for url in ("/admin", f"/admin/enroll/{uid}", "/teacher", "/teacher/reports", "/account/password"):
            self.assertEqual(admin.get(url).status_code, 200, url)
        t = self.client("pt", "pass1234")
        sid = self.start_session(t, "PAGES")
        self.assertEqual(t.get(f"/teacher/sessions/{sid}").status_code, 200)
        self.enroll(admin, "P1", 70)
        s = self.client("P1", "pass1234")
        for url in ("/student", "/student/scan"):
            self.assertEqual(s.get(url).status_code, 200, url)
        self.assertIn("Scan my face", s.get("/student").get_data(as_text=True))
        self.assertEqual(self.client().get("/login").status_code, 200)
        # starting a second session for the same section/day reuses the first
        self.assertEqual(self.start_session(t, "PAGES"), sid)

    def test_login_lockout_and_import(self):
        admin = self.client("admin", "adminpass")
        r = self.post(admin, "/admin/users/import", csv="Name,Roll,Section\nZed One,Z1,IMP\nZed Two,Z2,IMP\nbad line")
        body = r.get_data(as_text=True)
        self.assertIn("Zed One", body)
        c = self.client()
        for _ in range(6):
            r = c.post("/login", data={"username": "Z1", "password": "nope", "csrf": self.csrf(c)})
        self.assertIn("Too many", r.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main(verbosity=2)
