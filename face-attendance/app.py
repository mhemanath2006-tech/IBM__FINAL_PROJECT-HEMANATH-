import csv
import hmac
import io
import os
import secrets
import time
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (Flask, Response, abort, flash, g, jsonify, redirect,
                   render_template, request, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

import face_engine as fe
from db import IntegrityErrors, close_db, get_db, init_db

# ---------------------------------------------------------------- config
TZ = timezone(timedelta(minutes=int(os.environ.get("TZ_OFFSET_MINUTES", "330"))))  # IST
FMT = "%Y-%m-%d %H:%M:%S"
MATCH_THRESHOLD = float(os.environ.get("MATCH_THRESHOLD", "0.45"))  # ArcFace cosine
YAW_DELTA = float(os.environ.get("YAW_DELTA", "0.12"))  # required head-turn amount
LOW_PCT = float(os.environ.get("LOW_ATTENDANCE_PCT", "75"))
CHALLENGE_SECONDS = 90
LIVENESS_DEBUG = os.environ.get("LIVENESS_DEBUG") == "1"
ROLES = ("admin", "teacher", "student")


def create_app(testing=False):
    app = Flask(__name__)
    secret = os.environ.get("SECRET_KEY")
    if not secret:
        secret = secrets.token_hex(32)
        print("WARNING: SECRET_KEY not set; logins reset on every restart.", flush=True)
    app.config.update(
        SECRET_KEY=secret,
        MAX_CONTENT_LENGTH=8 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=not testing and os.environ.get("INSECURE_COOKIES") != "1",
        TESTING=testing,
    )
    app.teardown_appcontext(close_db)
    with app.app_context():
        init_db()
        bootstrap_admin()
    register(app)
    return app


def now():
    return datetime.now(TZ).replace(tzinfo=None)


def ts(dt):
    return dt.strftime(FMT)


def bootstrap_admin():
    from db import _connect

    conn = _connect()
    try:
        if conn.one("SELECT id FROM users WHERE role = 'admin'"):
            return
        username = os.environ.get("ADMIN_USERNAME", "admin")
        password = os.environ.get("ADMIN_PASSWORD")
        generated = not password
        if generated:
            password = secrets.token_urlsafe(9)
        conn.insert(
            "INSERT INTO users (username, password_hash, role, name) VALUES (?,?,?,?)",
            (username, generate_password_hash(password), "admin", "Administrator"),
        )
        if generated:
            print(f"FIRST RUN: admin login is '{username}' / '{password}' - change it.", flush=True)
    finally:
        conn.close()


# ---------------------------------------------------------------- helpers
def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


def new_password():
    return secrets.token_urlsafe(6)


def session_state(s, t=None):
    """open / waiting / closed for a class_sessions row."""
    t = ts(t or now())
    if s["closed"] or t >= s["closes_at"]:
        return "closed"
    return "open" if t >= s["opens_at"] else "waiting"


def todays_session(section):
    return get_db().one(
        "SELECT * FROM class_sessions WHERE section = ? AND day = ? ORDER BY id DESC LIMIT 1",
        (section, now().strftime("%Y-%m-%d")),
    )


def summarize(section=None, start=None, end=None):
    """Per-student attendance summary. Sessions still open do not count yet."""
    db = get_db()
    students = db.all(
        "SELECT id, name, roll_no, section FROM users WHERE role = 'student' AND active = 1"
        + (" AND section = ?" if section else "") + " ORDER BY section, roll_no, name",
        [section] if section else [],
    )
    q, args = "SELECT * FROM class_sessions WHERE 1=1", []
    if section:
        q += " AND section = ?"; args.append(section)
    if start:
        q += " AND day >= ?"; args.append(start)
    if end:
        q += " AND day <= ?"; args.append(end)
    sessions = [s for s in db.all(q + " ORDER BY day", args) if session_state(s) == "closed"]
    att = {}
    if sessions:
        ids = ",".join(str(int(s["id"])) for s in sessions)
        for a in db.all(f"SELECT session_id, student_id, status FROM attendance WHERE session_id IN ({ids})"):
            att[(a["session_id"], a["student_id"])] = a["status"]
    rows = []
    for st in students:
        mine = [s for s in sessions if s["section"] == st["section"]]
        present = sum(1 for s in mine if att.get((s["id"], st["id"])) == "present")
        late = sum(1 for s in mine if att.get((s["id"], st["id"])) == "late")
        total = len(mine)
        pct = round(100.0 * (present + late) / total, 1) if total else None
        rows.append({**st, "total": total, "present": present, "late": late,
                     "absent": total - present - late, "pct": pct,
                     "low": pct is not None and pct < LOW_PCT})
    return rows, sessions, att


def record_attendance(session_row, student_id, status, method, similarity=None):
    """Insert or update one attendance row."""
    db = get_db()
    existing = db.one("SELECT id FROM attendance WHERE session_id = ? AND student_id = ?",
                      (session_row["id"], student_id))
    if existing:
        db.run("UPDATE attendance SET status=?, method=?, marked_at=?, similarity=? WHERE id=?",
               (status, method, ts(now()), similarity, existing["id"]))
        return
    try:
        db.insert("INSERT INTO attendance (session_id, student_id, status, method, marked_at, similarity)"
                  " VALUES (?,?,?,?,?,?)",
                  (session_row["id"], student_id, status, method, ts(now()), similarity))
    except IntegrityErrors:
        pass  # a concurrent request already marked it


# ---------------------------------------------------------------- login rate limit
_fails = {}


def locked_out(key):
    n, until = _fails.get(key, (0, 0))
    return n >= 5 and time.time() < until


def note_fail(key):
    n, _ = _fails.get(key, (0, 0))
    _fails[key] = (n + 1, time.time() + 300)


# ---------------------------------------------------------------- routes
def register(app):
    app.jinja_env.globals.update(csrf_token=csrf_token, LOW_PCT=LOW_PCT)

    def role_required(*roles):
        def deco(fn):
            @wraps(fn)
            def wrapper(*a, **kw):
                if not g.user:
                    if request.path.startswith("/api/"):
                        return jsonify(ok=False, error="Please log in again."), 401
                    return redirect(url_for("login"))
                if g.user["role"] not in roles:
                    if request.path.startswith("/api/"):
                        return jsonify(ok=False, error="Not allowed."), 403
                    abort(403)
                return fn(*a, **kw)
            return wrapper
        return deco

    @app.before_request
    def load_user():
        g.user = None
        uid = session.get("uid")
        if uid:
            u = get_db().one("SELECT * FROM users WHERE id = ? AND active = 1", (uid,))
            if u:
                g.user = u
            else:
                session.pop("uid", None)
        if request.method == "POST":
            sent = request.form.get("csrf") or request.headers.get("X-CSRF-Token") or ""
            if not hmac.compare_digest(sent, session.get("csrf", "x")):
                if request.path.startswith("/api/"):
                    return jsonify(ok=False, error="Session expired. Reload the page."), 400
                abort(400, "Bad or missing CSRF token. Reload the page and try again.")

    @app.after_request
    def headers(resp):
        resp.headers["X-Content-Type-Options"] = "nosniff"
        resp.headers["Referrer-Policy"] = "same-origin"
        resp.headers["Permissions-Policy"] = "camera=(self)"
        if g.get("user"):
            resp.headers["Cache-Control"] = "no-store"
        return resp

    @app.context_processor
    def inject():
        return {"user": g.get("user")}

    # ---- auth
    @app.get("/")
    def index():
        if not g.user:
            return redirect(url_for("login"))
        return redirect(url_for({"admin": "admin_home", "teacher": "teacher_home",
                                 "student": "student_home"}[g.user["role"]]))

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            key = (request.remote_addr, username.lower())
            if locked_out(key):
                flash("Too many failed attempts. Try again in a few minutes.", "error")
            else:
                u = get_db().one("SELECT * FROM users WHERE username = ? AND active = 1", (username,))
                if u and check_password_hash(u["password_hash"], request.form.get("password", "")):
                    _fails.pop(key, None)
                    session.clear()
                    session["uid"] = u["id"]
                    csrf_token()
                    return redirect(url_for("index"))
                note_fail(key)
                flash("Wrong username or password.", "error")
        return render_template("login.html")

    @app.post("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    @app.route("/account/password", methods=["GET", "POST"])
    @role_required(*ROLES)
    def change_password():
        if request.method == "POST":
            cur, new = request.form.get("current", ""), request.form.get("new", "")
            if not check_password_hash(g.user["password_hash"], cur):
                flash("Current password is wrong.", "error")
            elif len(new) < 6:
                flash("New password must be at least 6 characters.", "error")
            else:
                get_db().run("UPDATE users SET password_hash = ? WHERE id = ?",
                             (generate_password_hash(new), g.user["id"]))
                flash("Password changed.", "ok")
                return redirect(url_for("index"))
        return render_template("password.html")

    # ---- admin
    @app.get("/admin")
    @role_required("admin")
    def admin_home():
        users = get_db().all(
            "SELECT u.*, (SELECT COUNT(*) FROM face_templates f WHERE f.user_id = u.id) AS has_face"
            " FROM users u ORDER BY u.role, u.section, u.roll_no, u.name")
        return render_template("admin.html", users=users)

    @app.post("/admin/users")
    @role_required("admin")
    def admin_add_user():
        f = request.form
        role, name = f.get("role", ""), f.get("name", "").strip()
        roll, section = f.get("roll_no", "").strip(), f.get("section", "").strip()
        username = f.get("username", "").strip() or roll
        if role not in ("teacher", "student") or not name or not username:
            flash("Name, role and username (or roll number) are required.", "error")
            return redirect(url_for("admin_home"))
        if role == "student" and not section:
            flash("Students need a section/class.", "error")
            return redirect(url_for("admin_home"))
        pw = f.get("password", "").strip() or new_password()
        try:
            get_db().insert(
                "INSERT INTO users (username, password_hash, role, name, roll_no, section) VALUES (?,?,?,?,?,?)",
                (username, generate_password_hash(pw), role, name, roll or None, section or None))
        except IntegrityErrors:
            flash(f"Username '{username}' already exists.", "error")
            return redirect(url_for("admin_home"))
        return render_template("credentials.html", creds=[(name, username, pw)])

    @app.post("/admin/users/import")
    @role_required("admin")
    def admin_import():
        """CSV lines: name,roll_no,section  -> student accounts with random passwords."""
        creds, skipped = [], []
        for row in csv.reader(io.StringIO(request.form.get("csv", ""))):
            row = [c.strip() for c in row]
            if len(row) < 3 or not all(row[:3]) or row[0].lower() == "name":
                if any(row):
                    skipped.append(",".join(row))
                continue
            name, roll, section = row[:3]
            pw = new_password()
            try:
                get_db().insert(
                    "INSERT INTO users (username, password_hash, role, name, roll_no, section) VALUES (?,?,?,?,?,?)",
                    (roll, generate_password_hash(pw), "student", name, roll, section))
                creds.append((name, roll, pw))
            except IntegrityErrors:
                skipped.append(f"{roll} (already exists)")
        return render_template("credentials.html", creds=creds, skipped=skipped)

    @app.post("/admin/users/<int:uid>/reset")
    @role_required("admin")
    def admin_reset(uid):
        u = get_db().one("SELECT * FROM users WHERE id = ?", (uid,)) or abort(404)
        pw = new_password()
        get_db().run("UPDATE users SET password_hash = ? WHERE id = ?", (generate_password_hash(pw), uid))
        return render_template("credentials.html", creds=[(u["name"], u["username"], pw)])

    @app.post("/admin/users/<int:uid>/toggle")
    @role_required("admin")
    def admin_toggle(uid):
        if uid == g.user["id"]:
            flash("You cannot deactivate yourself.", "error")
        else:
            get_db().run("UPDATE users SET active = 1 - active WHERE id = ?", (uid,))
        return redirect(url_for("admin_home"))

    @app.post("/admin/users/<int:uid>/face/delete")
    @role_required("admin")
    def admin_face_delete(uid):
        get_db().run("DELETE FROM face_templates WHERE user_id = ?", (uid,))
        get_db().run("UPDATE users SET consent_at = NULL WHERE id = ?", (uid,))
        flash("Face data deleted.", "ok")
        return redirect(url_for("admin_home"))

    @app.get("/admin/enroll/<int:uid>")
    @role_required("admin")
    def admin_enroll(uid):
        st = get_db().one("SELECT * FROM users WHERE id = ? AND role = 'student'", (uid,)) or abort(404)
        return render_template("enroll.html", student=st)

    @app.post("/api/admin/enroll/<int:uid>")
    @role_required("admin")
    def api_enroll(uid):
        st = get_db().one("SELECT * FROM users WHERE id = ? AND role = 'student'", (uid,))
        if not st:
            return jsonify(ok=False, error="Student not found."), 404
        data = request.get_json(silent=True) or {}
        if not data.get("consent"):
            return jsonify(ok=False, error="Tick the consent box first."), 400
        frames = data.get("frames") or []
        if not 3 <= len(frames) <= 10:
            return jsonify(ok=False, error="Send between 3 and 10 frames."), 400
        engine, good, problems = fe.get_engine(), [], []
        for i, fr in enumerate(frames, 1):
            try:
                good.append(fe.single_face(engine, fe.decode_image(fr))["embedding"])
            except fe.FaceError as e:
                problems.append(f"frame {i}: {e}")
        if len(good) < 3:
            return jsonify(ok=False, error="Not enough clear frames. " + "; ".join(problems[:2])), 400
        ref = good[0]
        if min(float(ref @ e) for e in good[1:]) < 0.5:
            return jsonify(ok=False, error="The frames do not look like the same person. Try again."), 400
        db = get_db()
        db.run("DELETE FROM face_templates WHERE user_id = ?", (uid,))
        db.run("INSERT INTO face_templates (user_id, embeddings, created_at) VALUES (?,?,?)",
               (uid, fe.encode_templates(good), ts(now())))
        db.run("UPDATE users SET consent_at = ? WHERE id = ?", (ts(now()), uid))
        return jsonify(ok=True, saved=len(good))

    # ---- teacher
    @app.get("/teacher")
    @role_required("teacher", "admin")
    def teacher_home():
        db = get_db()
        sections = [r["section"] for r in db.all(
            "SELECT DISTINCT section FROM users WHERE role='student' AND active=1 AND section IS NOT NULL ORDER BY section")]
        recent = db.all("SELECT * FROM class_sessions ORDER BY day DESC, id DESC LIMIT 15")
        for s in recent:
            s["state"] = session_state(s)
            s["count"] = db.one("SELECT COUNT(*) AS n FROM attendance WHERE session_id = ?", (s["id"],))["n"]
        return render_template("teacher.html", sections=sections, sessions=recent,
                               today=now().strftime("%Y-%m-%d"))

    @app.post("/teacher/sessions")
    @role_required("teacher", "admin")
    def teacher_start():
        section = request.form.get("section", "").strip()
        if not section:
            flash("Choose a section.", "error")
            return redirect(url_for("teacher_home"))
        existing = todays_session(section)
        if existing:
            flash("Today's session for this section already exists.", "ok")
            return redirect(url_for("teacher_session", sid=existing["id"]))
        try:
            open_for = max(5, min(int(request.form.get("open_minutes", 30)), 240))
            late_after = max(1, min(int(request.form.get("late_minutes", 10)), open_for))
        except ValueError:
            flash("Minutes must be numbers.", "error")
            return redirect(url_for("teacher_home"))
        t = now()
        sid = get_db().insert(
            "INSERT INTO class_sessions (section, teacher_id, day, opens_at, late_after, closes_at)"
            " VALUES (?,?,?,?,?,?)",
            (section, g.user["id"], t.strftime("%Y-%m-%d"), ts(t),
             ts(t + timedelta(minutes=late_after)), ts(t + timedelta(minutes=open_for))))
        return redirect(url_for("teacher_session", sid=sid))

    @app.get("/teacher/sessions/<int:sid>")
    @role_required("teacher", "admin")
    def teacher_session(sid):
        db = get_db()
        s = db.one("SELECT * FROM class_sessions WHERE id = ?", (sid,)) or abort(404)
        roster = db.all(
            "SELECT u.id, u.name, u.roll_no, a.status, a.method, a.marked_at"
            " FROM users u LEFT JOIN attendance a ON a.student_id = u.id AND a.session_id = ?"
            " WHERE u.role='student' AND u.active=1 AND u.section = ? ORDER BY u.roll_no, u.name",
            (sid, s["section"]))
        counts = {k: sum(1 for r in roster if (r["status"] or "absent") == k) for k in ("present", "late", "absent")}
        return render_template("session.html", s=s, state=session_state(s), roster=roster, counts=counts)

    @app.post("/teacher/sessions/<int:sid>/mark")
    @role_required("teacher", "admin")
    def teacher_mark(sid):
        db = get_db()
        s = db.one("SELECT * FROM class_sessions WHERE id = ?", (sid,)) or abort(404)
        status = request.form.get("status")
        try:
            student_id = int(request.form.get("student_id", ""))
        except ValueError:
            abort(400)
        st = db.one("SELECT id FROM users WHERE id = ? AND role='student' AND section = ?",
                    (student_id, s["section"]))
        if status not in ("present", "late", "absent") or not st:
            abort(400)
        record_attendance(s, student_id, status, "manual")
        return redirect(url_for("teacher_session", sid=sid))

    @app.post("/teacher/sessions/<int:sid>/close")
    @role_required("teacher", "admin")
    def teacher_close(sid):
        get_db().run("UPDATE class_sessions SET closed = 1 WHERE id = ?", (sid,))
        return redirect(url_for("teacher_session", sid=sid))

    @app.post("/teacher/sessions/<int:sid>/extend")
    @role_required("teacher", "admin")
    def teacher_extend(sid):
        db = get_db()
        s = db.one("SELECT * FROM class_sessions WHERE id = ?", (sid,)) or abort(404)
        new_close = ts(now() + timedelta(minutes=10))
        db.run("UPDATE class_sessions SET closed = 0, closes_at = ? WHERE id = ?", (new_close, sid))
        return redirect(url_for("teacher_session", sid=sid))

    def report_args():
        return (request.args.get("section", "").strip() or None,
                request.args.get("from", "").strip() or None,
                request.args.get("to", "").strip() or None)

    @app.get("/teacher/reports")
    @role_required("teacher", "admin")
    def teacher_reports():
        section, start, end = report_args()
        rows, sessions, _ = summarize(section, start, end)
        sections = [r["section"] for r in get_db().all(
            "SELECT DISTINCT section FROM users WHERE role='student' AND section IS NOT NULL ORDER BY section")]
        return render_template("reports.html", rows=rows, n_sessions=len(sessions), sections=sections,
                               section=section or "", start=start or "", end=end or "")

    @app.get("/teacher/export.xlsx")
    @role_required("teacher", "admin")
    def teacher_export():
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        section, start, end = report_args()
        rows, sessions, att = summarize(section, start, end)
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Roll no", "Name", "Section", "Classes", "Present", "Late", "Absent", "Attendance %"])
        for r in rows:
            ws.append([r["roll_no"], r["name"], r["section"], r["total"], r["present"], r["late"],
                       r["absent"], r["pct"]])
            if r["low"]:
                for c in ws[ws.max_row]:
                    c.fill = PatternFill("solid", fgColor="FFD9D9")
        for c in ws[1]:
            c.font = Font(bold=True)
        ws2 = wb.create_sheet("Daily")
        ws2.append(["Roll no", "Name"] + [f'{s["day"]} ({s["section"]})' for s in sessions])
        for r in rows:
            line = [r["roll_no"], r["name"]]
            for s in sessions:
                if s["section"] != r["section"]:
                    line.append("")
                else:
                    line.append((att.get((s["id"], r["id"]), "absent")).capitalize())
            ws2.append(line)
        for c in ws2[1]:
            c.font = Font(bold=True)
        for sheet in (ws, ws2):
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = 16
        buf = io.BytesIO()
        wb.save(buf)
        return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": "attachment; filename=attendance.xlsx"})

    # ---- student
    @app.get("/student")
    @role_required("student")
    def student_home():
        db = get_db()
        u = g.user
        s = todays_session(u["section"]) if u["section"] else None
        mine = db.one("SELECT * FROM attendance WHERE session_id = ? AND student_id = ?",
                      (s["id"], u["id"])) if s else None
        has_face = bool(db.one("SELECT user_id FROM face_templates WHERE user_id = ?", (u["id"],)))
        rows, _, _ = summarize(u["section"])
        me = next((r for r in rows if r["id"] == u["id"]), None)
        history = db.all(
            "SELECT c.day, a.status, a.method, a.marked_at FROM attendance a"
            " JOIN class_sessions c ON c.id = a.session_id WHERE a.student_id = ? ORDER BY c.day DESC LIMIT 30",
            (u["id"],))
        return render_template("student.html", s=s, state=session_state(s) if s else None,
                               mine=mine, has_face=has_face, me=me, history=history)

    @app.get("/student/scan")
    @role_required("student")
    def student_scan():
        return render_template("scan.html")

    @app.post("/api/scan/challenge")
    @role_required("student")
    def api_challenge():
        db, u = get_db(), g.user
        if not u["section"]:
            return jsonify(ok=False, error="You have no section assigned. Ask the admin."), 400
        if not db.one("SELECT user_id FROM face_templates WHERE user_id = ?", (u["id"],)):
            return jsonify(ok=False, error="Your face is not registered yet. Ask the admin."), 400
        s = todays_session(u["section"])
        if not s or session_state(s) != "open":
            return jsonify(ok=False, error="Attendance is not open right now."), 400
        if db.one("SELECT id FROM attendance WHERE session_id = ? AND student_id = ?", (s["id"], u["id"])):
            return jsonify(ok=False, error="Your attendance is already marked."), 400
        recent = db.one("SELECT COUNT(*) AS n FROM challenges WHERE user_id = ? AND created_at >= ?",
                        (u["id"], ts(now() - timedelta(minutes=10))))["n"]
        if recent >= 10:
            return jsonify(ok=False, error="Too many attempts. Wait a few minutes or ask your teacher."), 429
        nonce = secrets.token_urlsafe(16)
        direction = secrets.choice(["left", "right"])
        db.run("INSERT INTO challenges (nonce, user_id, direction, created_at, expires_at) VALUES (?,?,?,?,?)",
               (nonce, u["id"], direction, ts(now()), ts(now() + timedelta(seconds=CHALLENGE_SECONDS))))
        return jsonify(ok=True, nonce=nonce, direction=direction)

    @app.post("/api/scan/submit")
    @role_required("student")
    def api_submit():
        db, u = get_db(), g.user
        data = request.get_json(silent=True) or {}
        neutral, turns = data.get("neutral"), data.get("turns") or []
        ch = db.one("SELECT * FROM challenges WHERE nonce = ? AND user_id = ?", (data.get("nonce", ""), u["id"]))
        if not ch or ch["used"] or ts(now()) > ch["expires_at"]:
            return jsonify(ok=False, error="That scan expired. Start again."), 400
        if db.run("UPDATE challenges SET used = 1 WHERE nonce = ? AND used = 0", (ch["nonce"],)) != 1:
            return jsonify(ok=False, error="That scan was already used."), 400
        s = todays_session(u["section"]) if u["section"] else None
        if not s or session_state(s) != "open":
            return jsonify(ok=False, error="Attendance closed before your scan finished."), 400
        if db.one("SELECT id FROM attendance WHERE session_id = ? AND student_id = ?", (s["id"], u["id"])):
            return jsonify(ok=False, error="Your attendance is already marked."), 400
        row = db.one("SELECT embeddings FROM face_templates WHERE user_id = ?", (u["id"],))
        if not row:
            return jsonify(ok=False, error="Your face is not registered yet."), 400
        if not isinstance(turns, list) or not 3 <= len(turns) <= 12 or not neutral:
            return jsonify(ok=False, error="Incomplete scan. Try again."), 400

        templates, engine = fe.decode_templates(row["embeddings"]), fe.get_engine()
        try:
            f0 = fe.single_face(engine, fe.decode_image(neutral))
        except fe.FaceError as e:
            return jsonify(ok=False, error=f"Straight-on photo: {e}"), 400
        sim0 = fe.best_similarity(f0["embedding"], templates)
        if sim0 < MATCH_THRESHOLD:
            return jsonify(ok=False, error="Your face did not match your registered face."), 400
        yaw0, want, best_sim, best_delta, matched = fe.yaw_ratio(f0["kps"]), ch["direction"], sim0, 0.0, False
        for fr in turns:
            try:
                f = fe.single_face(engine, fe.decode_image(fr))
            except fe.FaceError:
                continue  # head turned too far or blurry frame - skip it
            sim = fe.best_similarity(f["embedding"], templates)
            delta = fe.yaw_ratio(f["kps"]) - yaw0
            signed = delta if want == "left" else -delta
            best_delta = max(best_delta, signed)
            if sim >= MATCH_THRESHOLD and signed >= YAW_DELTA:
                matched, best_sim = True, max(best_sim, sim)
        if not matched:
            out = {"ok": False, "error": f"Liveness check failed. Turn your head to your {want} when asked."}
            if LIVENESS_DEBUG:
                out["debug"] = {"best_delta": best_delta, "need": YAW_DELTA, "sim0": sim0}
            return jsonify(out), 400
        status = "present" if ts(now()) <= s["late_after"] else "late"
        record_attendance(s, u["id"], status, "face", best_sim)
        return jsonify(ok=True, status=status, time=now().strftime("%H:%M"))
